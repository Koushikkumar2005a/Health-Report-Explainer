"""
Health Report Explainer - Document Processing Service

Handles text extraction from various document formats:
- PDF (text-based and scanned)
- DOCX
- Images (JPG, JPEG, PNG) via Gemini Vision
"""
import os
import logging
from pypdf import PdfReader
from docx import Document
from PIL import Image
from google import genai
from backend.config import settings

MODEL_NAME = settings.GEMINI_MODEL

logger = logging.getLogger(__name__)


class DocumentService:
    """Service for extracting text from uploaded documents."""

    def __init__(self):
        self.client = None
        if settings.GEMINI_API_KEY:
            self.client = genai.Client(api_key=settings.GEMINI_API_KEY)

    async def extract_text(self, file_path: str, filename: str) -> str:
        """Extract text from a document based on its file type."""
        ext = os.path.splitext(filename)[1].lower()

        try:
            if ext in (".jpg", ".jpeg", ".png"):
                return await self._extract_from_image(file_path)
            elif ext == ".pdf":
                return await self._extract_from_pdf(file_path)
            elif ext == ".doc":
                raise ValueError(
                    "Legacy DOC files cannot be processed reliably in this environment. "
                    "Please convert the file to DOCX or PDF before uploading."
                )
            elif ext == ".docx":
                return await self._extract_from_docx(file_path)
            else:
                raise ValueError(f"Unsupported file type: {ext}")
        except Exception as e:
            logger.error(f"Text extraction failed for {filename}: {str(e)}")
            raise

    async def _extract_from_image(self, file_path: str) -> str:
        """Extract text from image using Gemini Vision."""
        try:
            img = Image.open(file_path)
            if img.mode in ("RGBA", "P"):
                img = img.convert("RGB")

            if not self.client:
                raise ValueError("Gemini API is not configured for OCR.")

            response = self.client.models.generate_content(
                model=MODEL_NAME,
                contents=[
                    "Extract ALL text from this medical/health report image. "
                    "Include every value, label, reference range, unit, patient detail, "
                    "test name, result, and any other text visible in the image. "
                    "Maintain the structure and organization of the data as closely as possible. "
                    "If there are tables, preserve the table structure. "
                    "Do not interpret or explain the results — just extract the raw text faithfully.",
                    img,
                ],
            )

            text = response.text.strip()
            if not text:
                raise ValueError("No text could be extracted from the image.")
            return text

        except Exception as e:
            logger.exception("Image OCR failed for file=%s", file_path)
            raise ValueError(
                "We could not read the uploaded image. "
                "Please upload a clearer image of your report."
            ) from e

    async def _extract_from_pdf(self, file_path: str) -> str:
        """Extract text from PDF. Falls back to vision for scanned PDFs."""
        try:
            reader = PdfReader(file_path)
            text_parts = []
            has_text = False

            for page_num, page in enumerate(reader.pages):
                text = page.extract_text()
                if text and text.strip():
                    has_text = True
                    text_parts.append(f"--- Page {page_num + 1} ---\n{text.strip()}")

            if has_text and len("\n".join(text_parts)) > 50:
                return "\n\n".join(text_parts)

            # Scanned PDF — use Gemini Vision on the file directly
            logger.info("PDF appears scanned, using vision extraction...")
            return await self._extract_scanned_pdf(file_path)

        except ValueError:
            raise
        except Exception as e:
            logger.error(f"PDF extraction failed: {str(e)}")
            raise ValueError(
                "We could not read the uploaded PDF. "
                "Please ensure it is a valid PDF document."
            ) from e

    async def _extract_scanned_pdf(self, file_path: str) -> str:
        """Extract text from scanned PDF by sending it to Gemini Vision."""
        try:
            if not self.client:
                raise ValueError("Gemini API is not configured for OCR.")

            uploaded_file = self.client.files.upload(
                file=file_path,
                config={"mime_type": "application/pdf"},
            )

            try:
                response = self.client.models.generate_content(
                    model=MODEL_NAME,
                    contents=[
                        "Extract ALL text from this medical/health report PDF. "
                        "Include every value, label, reference range, unit, patient detail, "
                        "test name, result, and any other text visible in the document. "
                        "Maintain the structure and organization of the data as closely as possible. "
                        "If there are tables, preserve the table structure. "
                        "Do not interpret or explain the results — just extract the raw text faithfully.",
                        uploaded_file,
                    ],
                )
            finally:
                try:
                    self.client.files.delete(name=uploaded_file.name)
                except Exception:
                    pass

            text = response.text.strip()
            if not text:
                raise ValueError("No text could be extracted from the scanned PDF.")
            return text

        except ValueError:
            raise
        except Exception as e:
            logger.exception("Scanned PDF extraction failed for file=%s", file_path)
            raise ValueError(
                "We could not read the scanned PDF. "
                "Please upload a clearer version of your report."
            ) from e

    async def _extract_from_docx(self, file_path: str) -> str:
        """Extract text from DOCX files."""
        try:
            doc = Document(file_path)
            text_parts = []

            # Extract paragraphs
            for para in doc.paragraphs:
                if para.text.strip():
                    text_parts.append(para.text.strip())

            # Extract tables
            for table in doc.tables:
                table_rows = []
                for row in table.rows:
                    cells = [cell.text.strip() for cell in row.cells]
                    table_rows.append(" | ".join(cells))
                if table_rows:
                    text_parts.append("\n".join(table_rows))

            text = "\n\n".join(text_parts)
            if not text.strip():
                raise ValueError("The document appears to be empty.")

            return text

        except ValueError:
            raise
        except Exception as e:
            logger.error(f"DOCX extraction failed: {str(e)}")
            raise ValueError(
                "We could not read the uploaded Word document. "
                "Please ensure it is a valid DOCX file."
            ) from e


# Singleton
document_service = DocumentService()
